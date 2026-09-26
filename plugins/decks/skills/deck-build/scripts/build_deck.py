#!/usr/bin/env python3
"""
build_deck.py — build a .pptx from a template's OWN layouts and placeholders,
driven by a plain outline file. No drawing, no hex colours, no free text boxes:
every slide is "add layout X, fill placeholder Y" so the deck stays a real
template document that the owner can restyle globally in PowerPoint.

  python3 build_deck.py outline.yaml
  python3 build_deck.py outline.yaml --out Deck.pptx --template T.potx --master "Indigo · Light"

Outline (YAML or JSON) — see references/outline.md in this skill:

  template: path/to/Template.potx        # or --template
  master: Luminous · Indigo · Light      # default master (name or unique substring)
  out: path/to/Deck.pptx                 # or --out
  defaults:                              # applied wherever the placeholder exists
    Epistemic label: I know this
  slides:
    - layout: Cover
      Eyebrow: AI foundations · Free series
      Title: Stay the one who decides
      Subtitle: A short course on using AI tools without outsourcing your judgment.
      notes: Speaker notes go here.
    - layout: Content
      Title: Three zones
      Body:
        - Green — use freely
        - Yellow — use with a rule
        - {text: A sub-point, level: 1}
      Epistemic label: I practice this
    - layout: Image right
      master: Luminous · Rose · Light     # per-slide master override
      Picture: images/owl.jpg

Keys are placeholder NAMES (case-insensitive; spaces/underscores/dashes ignored)
or `idx:N`. Unknown keys are errors, so typos never silently vanish. A string
becomes paragraphs split on newlines; a list becomes one paragraph per item;
`**bold**` inside text becomes a bold run. Placeholders you do not fill are
removed from the slide (set `keep_empty: true` on a slide to keep them).

Needs: python-pptx (+ pyyaml for YAML outlines).
"""
from __future__ import annotations
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

POTX_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
PPTX_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
RESERVED = {"layout", "master", "notes", "keep_empty"}


# --------------------------------------------------------------------------- #
def as_pptx(path: str) -> str:
    """python-pptx refuses .potx by content type — re-type a temp copy."""
    if not path.lower().endswith(".potx"):
        return path
    tmp = tempfile.mkdtemp(prefix="deck-")
    out = os.path.join(tmp, os.path.basename(path)[:-5] + ".pptx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(POTX_CT.encode(), PPTX_CT.encode())
            zout.writestr(item, data)
    return out


def norm(s: str) -> str:
    return re.sub(r"[\s_\-·]+", "", str(s)).lower()


def load_outline(path: str) -> dict:
    with open(path, encoding="utf-8") as f:
        text = f.read()
    if path.lower().endswith(".json"):
        return json.loads(text)
    try:
        import yaml
    except ImportError:
        sys.exit("Outline is YAML but PyYAML is not installed: pip install pyyaml (or use JSON).")
    return yaml.safe_load(text)


def master_name(m) -> str:
    try:
        from lxml import etree
        part = m.part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
        return etree.fromstring(part.blob).get("name") or m.name or ""
    except Exception:
        return m.name or ""


def pick(items, wanted: str, what: str, names):
    """Exact (normalised) match first, then unique substring match."""
    if wanted is None:
        return items[0]
    w = norm(wanted)
    exact = [it for it, n in zip(items, names) if norm(n) == w]
    if len(exact) == 1:
        return exact[0]
    sub = [it for it, n in zip(items, names) if w in norm(n)]
    if len(sub) == 1:
        return sub[0]
    opts = "\n  ".join(names)
    if not sub:
        sys.exit(f"No {what} matches '{wanted}'. Options:\n  {opts}")
    sys.exit(f"'{wanted}' matches several {what}s — be more specific. Options:\n  {opts}")


def ph_kind(ph) -> str:
    return str(ph.placeholder_format.type).split(".")[-1].split(" ")[0].lower()


# --------------------------------------------------------------------------- #
BOLD = re.compile(r"\*\*(.+?)\*\*")


def write_runs(paragraph, text: str):
    """Plain text with optional **bold** spans → runs. Nothing else is styled:
    size, face and colour all inherit from the layout."""
    pos = 0
    for m in BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run().text = text[pos:m.start()]
        r = paragraph.add_run(); r.text = m.group(1); r.font.bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run().text = text[pos:]


def fill_text(ph, value):
    items = value if isinstance(value, list) else str(value).split("\n")
    tf = ph.text_frame
    first = True
    for it in items:
        level = 0
        if isinstance(it, dict):
            level = int(it.get("level", 0)); it = it.get("text", "")
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if level:
            p.level = level
        write_runs(p, str(it))


def fill_picture(ph, value, base_dir):
    path = value if os.path.isabs(value) else os.path.join(base_dir, value)
    if not os.path.exists(path):
        alt = os.path.abspath(value)
        if os.path.exists(alt):
            path = alt
        else:
            sys.exit(f"Picture not found: {value}")
    kind = ph_kind(ph)
    if kind == "picture":
        ph.insert_picture(path)
    else:  # object / body slot handed an image: place it inside the slot's bounds
        x, y, w, h = ph.left, ph.top, ph.width, ph.height
        slide = ph.part.slide
        ph._element.getparent().remove(ph._element)
        pic = slide.shapes.add_picture(path, x, y)
        # fit inside the box, keep aspect
        sc = min(w / pic.width, h / pic.height)
        pic.width = int(pic.width * sc); pic.height = int(pic.height * sc)
        pic.left = x + (w - pic.width) // 2; pic.top = y + (h - pic.height) // 2


IMAGE_EXT = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp", ".emf", ".svg")


def build(outline_path: str, out=None, template=None, master=None) -> str:
    from pptx import Presentation
    o = load_outline(outline_path)
    base = os.path.dirname(os.path.abspath(outline_path))
    template = template or o.get("template")
    if not template:
        sys.exit("No template: set `template:` in the outline or pass --template.")
    tpath = template if os.path.isabs(template) else os.path.join(base, template)
    if not os.path.exists(tpath):
        tpath = os.path.abspath(template)
    if not os.path.exists(tpath):
        sys.exit(f"Template not found: {template}")
    out = out or o.get("out") or os.path.splitext(outline_path)[0] + ".pptx"
    if not os.path.isabs(out):
        out = os.path.join(base, out)

    prs = Presentation(as_pptx(tpath))
    # A .pptx used as a template may carry slides — start clean.
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)

    masters = list(prs.slide_masters)
    mnames = [master_name(m) for m in masters]
    default_master = pick(masters, master or o.get("master"), "master", mnames)
    defaults = o.get("defaults") or {}
    warnings = []

    for n, sd in enumerate(o.get("slides") or [], 1):
        if not isinstance(sd, dict) or "layout" not in sd:
            sys.exit(f"Slide {n}: every slide needs a `layout:`.")
        m = pick(masters, sd["master"], "master", mnames) if sd.get("master") else default_master
        layouts = list(m.slide_layouts)
        layout = pick(layouts, sd["layout"], "layout", [l.name for l in layouts])
        slide = prs.slides.add_slide(layout)

        # python-pptx clones placeholders with generic names ("Text Placeholder 1");
        # the meaningful names live on the LAYOUT — resolve by idx and restore them.
        lay_names = {lp.placeholder_format.idx: lp.name for lp in layout.placeholders}
        phs = list(slide.placeholders)
        by_key = {}
        for ph in phs:
            idx = ph.placeholder_format.idx
            if idx in lay_names:
                ph.name = lay_names[idx]
            by_key[norm(ph.name)] = ph
            by_key[f"idx:{idx}"] = ph
        values = {k: v for k, v in defaults.items()}
        values.update({k: v for k, v in sd.items() if k not in RESERVED})
        filled = set()
        for key, val in values.items():
            k = key if str(key).startswith("idx:") else norm(key)
            ph = by_key.get(k)
            if ph is None:
                if key in defaults and key not in sd:
                    continue  # a default that this layout simply doesn't have
                names = ", ".join(sorted({p.name for p in phs if ph_kind(p) != "slide_number"}))
                sys.exit(f"Slide {n} ({layout.name}): no placeholder named '{key}'. "
                         f"This layout has: {names}")
            if val is None or (isinstance(val, str) and not val.strip()):
                continue
            kind = ph_kind(ph)
            ph_idx = ph.placeholder_format.idx  # read before fill: insert_picture swaps the element
            if kind == "picture" or (isinstance(val, str) and val.lower().endswith(IMAGE_EXT)):
                fill_picture(ph, str(val), base)
            elif kind in ("object",) and not isinstance(val, (str, list)):
                warnings.append(f"slide {n}: object slot '{ph.name}' given a non-text value; skipped")
                continue
            else:
                fill_text(ph, val)
            filled.add(ph_idx)

        if sd.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(sd["notes"])

        if not sd.get("keep_empty"):
            for ph in list(slide.placeholders):  # fresh: filled pictures replaced their elements
                if ph.placeholder_format.idx in filled or ph_kind(ph) == "slide_number":
                    continue
                el = ph._element
                if el.getparent() is not None:
                    el.getparent().remove(el)

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    n = len(prs.slides._sldIdLst)
    print(f"Built {out}  ({n} slides · master: {master_name(default_master)} · template: {os.path.basename(tpath)})")
    for w in warnings:
        print("  warning:", w)
    return out


def main(argv):
    if len(argv) < 2 or argv[1] in ("-h", "--help"):
        print(__doc__); return 0
    args = argv[1:]
    def opt(name):
        if name in args:
            i = args.index(name); v = args[i + 1]; del args[i:i + 2]; return v
    out, template, master = opt("--out"), opt("--template"), opt("--master")
    build(args[0], out=out, template=template, master=master)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
