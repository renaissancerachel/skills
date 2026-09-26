"""Shared drawing tools for Framework-slide diagrams (deck-build family).

A lesson's diagrams script imports these, defines one function per diagram
(fn(slide, x, y, w, h) drawing inside the Diagram slot), maps slide Titles to
functions, and calls draw(deck_path, DIAGRAMS) after build_deck.py.
Everything is drawn in theme colours and theme fonts, so it restyles with the
master; shapes carry no theme style, so no shadows.
"""
import os

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR as C
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HEAD, BODY = "+mj-lt", "+mn-lt"  # theme major (heading) / minor (body) font


# ── primitives ────────────────────────────────────────────────────────────
def _alpha(shape, val):
    fill = shape._element.spPr.find(qn("a:solidFill"))
    if fill is not None:
        clr = fill[0]
        a = etree.SubElement(clr, qn("a:alpha")); a.set("val", str(val))


def _unstyle(shape):
    """Drop the theme's default shape style (it adds a soft shadow to shape and text);
    fill, line, and font colours are all set explicitly from theme slots."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def box(slide, x, y, w, h, text="", style="fill", size=16, head=False, sub=None, sub_size=12,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER):
    """style: fill (Accent 1, light text) · soft (Accent 1 tint, dark text) · outline · plain."""
    s = slide.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    s.shadow.inherit = False
    _unstyle(s)
    ink = C.TEXT_1
    if style == "fill":
        s.fill.solid(); s.fill.fore_color.theme_color = C.ACCENT_1; _alpha(s, 90000)
        s.line.fill.background(); ink = C.BACKGROUND_1
    elif style == "soft":
        s.fill.solid(); s.fill.fore_color.theme_color = C.ACCENT_1
        s.fill.fore_color.brightness = 0.78; s.line.fill.background()
    elif style == "outline":
        s.fill.background(); s.line.color.theme_color = C.ACCENT_1; s.line.width = Pt(1.75)
    else:  # plain
        s.fill.background(); s.line.fill.background()
    tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Emu(91440 // 2))
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.size = Pt(size)
    r.font.name = HEAD if head else BODY; r.font.color.theme_color = ink
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = align
        r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(sub_size); r2.font.name = BODY
        r2.font.color.theme_color = ink
    return s


def arrow(slide, x1, y1, x2, y2, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(int(x1)), Emu(int(y1)),
                                   Emu(int(x2)), Emu(int(y2)))
    _unstyle(c)
    c.line.color.theme_color = C.ACCENT_1; c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    if dashed:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    return c


# ── runner ────────────────────────────────────────────────────────────────
def draw(path, diagrams):
    """Draw each Framework slide's diagram, matched by slide Title. If the outline
    dropped the slide's Caption, the diagram takes the caption area too."""
    prs = Presentation(path)
    done = []
    for n, slide in enumerate(prs.slides, 1):
        if slide.slide_layout.name != "Framework":
            continue
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title else ""
        fn = diagrams.get(title)
        if fn is None:
            print(f"  slide {n}: no diagram for '{title}'"); continue
        slot = next(p for p in slide.slide_layout.placeholders if "OBJECT" in str(p.placeholder_format.type))
        pad = int(slot.height * 0.04)
        height = slot.height
        has_caption = any(ph.name == "Caption" for ph in slide.placeholders)
        if not has_caption:  # caption dropped in the outline: the diagram takes its space too
            cap = next(p for p in slide.slide_layout.placeholders if p.name == "Caption")
            height = cap.top + cap.height - slot.top
        fn(slide, slot.left + pad, slot.top + pad, slot.width - 2 * pad, height - 2 * pad)
        done.append(n)
    prs.save(path)
    print(f"Diagrams drawn on slides {done} → {os.path.basename(path)}")

